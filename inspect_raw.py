from openpyxl import load_workbook
from pathlib import Path

files = [
    "SWaT_Dataset_Attack_v0.xlsx",
    "SWaT_Dataset_Normal_v0.xlsx",
    "SWaT_Dataset_Normal_v1.xlsx"
]

for name in files:

    path = Path("data/raw") / name

    print("\n" + "=" * 80)
    print(name)
    print("=" * 80)

    print("File size:",
          round(path.stat().st_size / (1024**2), 2),
          "MB")

    wb = load_workbook(
        filename=path,
        read_only=True,
        data_only=True
    )

    print("Sheets:", wb.sheetnames)

    for ws in wb.worksheets:

        print("\nSheet:", ws.title)
        print("Rows:", ws.max_row)
        print("Columns:", ws.max_column)

        for row in ws.iter_rows(
            min_row=1,
            max_row=3,
            values_only=True
        ):
            print(list(row)[:15])

    wb.close()
